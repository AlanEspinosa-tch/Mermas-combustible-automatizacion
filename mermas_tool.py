# -*- coding: utf-8 -*-
"""
HERRAMIENTA DE MERMAS Y CONCILIACIÓN DE PIPAS
=============================================
Entrada (carpeta o .zip, la estructura interna no importa):
  - Un Excel de pipas cuyo nombre contenga "PIPA"  (ej. "PIPAS DEL MES.xlsx")
  - Los CSV de mermas: "<ESTACION>- <Producto>.CSV"  (ej. "HUEHUETOCA- Premium 91 octanos.CSV")

Salida: un Excel con formato: Resumen, Mermas mensual, Compras sin pipa, Detalle diario,
        Conciliación pipas, Detalle pipas, Registros manuales, Parámetros y Notas.
        En la hoja "Registros manuales" del propio reporte se capturan las pipas que faltaban por
        registrar; todo el libro (mermas, semáforos, conciliación, resumen) se recalcula solo, sin volver a correr el script.

Uso local:   python mermas_tool.py  "C:/ruta/a/MERMAS"   (o un .zip)
Uso Colab:   ver cuaderno Mermas_Colab.ipynb
"""
import sys, os, re, zipfile, tempfile, unicodedata
from itertools import combinations
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, IconSetRule, CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

# ─────────────────────────── CONFIGURACIÓN ───────────────────────────
ESTACIONES = ["ATB", "ATP", "AVE FENIX", "AURORA", "COMBULUB", "HUEHUETOCA", "QL",
              "IXTAZACUALA", "TREHER", "TH", "TOREMEX", "TEOLOYUCAN"]
PRODUCTOS = ["Regular", "Premium", "Diesel"]
# Nombres alternos -> nombre oficial (lo que venga en PIPAS o en los CSV)
ALIAS = {"TREHER T": "TREHER", "ATEPO": "ATP", "AVE. FENIX": "AVE FENIX", "AVEFENIX": "AVE FENIX"}

LIMITE_SAT = 0.005        # 0.5 % de la venta
UMBRAL_ALERTA = 0.004     # 0.4 %  -> amarillo (cerca del límite)
VENTANA_DIAS = 2          # desfase permitido entre fecha BOL y registro en mermas
VENTANA_EXTENDIDA = 5     # se sigue buscando hasta aquí, pero se marca FUERA DE TOLERANCIA
TOLERANCIA_LTS = 5        # litros de tolerancia por pipa al conciliar
UMBRAL_CAPTURA = 1000    # compra y pipa cercanas con diferencia menor a esto => posible error de captura
UMBRAL_SALTO = 5000       # |merma diaria| >= esto => posible pipa registrada en otro día

P_LIM, P_ALE, P_TOL, P_SALTO = ("'Parámetros'!$C$5", "'Parámetros'!$C$6", "'Parámetros'!$C$8", "'Parámetros'!$C$9")

# ─────────────────────────── LECTURA ───────────────────────────
def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()

def estacion_oficial(s):
    s = norm(s)
    return ALIAS.get(s, s)

def producto_de(texto):
    t = norm(texto)
    if "DIESEL" in t: return "Diesel"
    if "PREMIUM" in t or "91" in t: return "Premium"
    if "REGULAR" in t or "MAGNA" in t or "87" in t: return "Regular"
    return None

def num(x):
    if pd.isna(x): return 0.0
    if isinstance(x, (int, float)): return float(x)
    return float(str(x).replace(",", "").replace("$", "").strip() or 0)

def leer_csv_mermas(path):
    for enc in ("utf-8-sig", "latin-1"):
        try:
            df = pd.read_csv(path, encoding=enc, dtype=str)
            break
        except UnicodeDecodeError:
            continue
    df.columns = [norm(c) for c in df.columns]
    col = {"FECHA": None, "SDO.INICIAL": None, "COMPRAS": None, "VENTAS": None,
           "AJUSTES": None, "SDO.FINAL": None, "SDO.REAL": None, "MERMA": None}
    for c in df.columns:
        for k in col:
            if col[k] is None and c.startswith(k) and not c.startswith("% "):
                col[k] = c
    total = df[df[col["FECHA"]].str.upper().str.contains("TOTAL", na=False)]
    df = df[~df[col["FECHA"]].str.upper().str.contains("TOTAL", na=False)].copy()
    out = pd.DataFrame({
        "Fecha": pd.to_datetime(df[col["FECHA"]], dayfirst=True, errors="coerce"),
        "SdoInicial": df[col["SDO.INICIAL"]].map(num),
        "Compras": df[col["COMPRAS"]].map(num),
        "Ventas": df[col["VENTAS"]].map(num),
        "Ajustes": df[col["AJUSTES"]].map(num) if col["AJUSTES"] else 0.0,
        "SdoFinal": df[col["SDO.FINAL"]].map(num),
        "SdoReal": df[col["SDO.REAL"]].map(num),
        "MermaArchivo": df[col["MERMA"]].map(num),
    }).dropna(subset=["Fecha"]).sort_values("Fecha")
    tot_merma = num(total[col["MERMA"]].iloc[0]) if len(total) else None
    return out, tot_merma

def localizar_entradas(origen):
    if zipfile.is_zipfile(origen):
        tmp = tempfile.mkdtemp()
        with zipfile.ZipFile(origen) as z:
            z.extractall(tmp)
        origen = tmp
    pipas, csvs = None, []
    for root, _, files in os.walk(origen):
        if "__MACOSX" in root: continue
        for f in files:
            p = os.path.join(root, f)
            if f.startswith("~$") or f.startswith("._"): continue
            if f.lower().endswith((".xlsx", ".xls")) and "PIPA" in norm(f):
                pipas = p
            elif f.lower().endswith(".csv"):
                csvs.append(p)
    return pipas, csvs

def leer_pipas(path):
    d = pd.read_excel(path)
    d.columns = [str(c).strip() for c in d.columns]
    prod_col = "COMBUSTIBLE2" if "COMBUSTIBLE2" in d.columns else "COMBUSTIBLE"
    p = pd.DataFrame({
        "Estacion": d["EMPRESA"].map(estacion_oficial),
        "Producto": d[prod_col].map(producto_de),
        "Factura": d.get("FACTURA"),
        "BOL": d.get("BOL"),
        "FechaBOL": pd.to_datetime(d["FECHA BOL"], errors="coerce").dt.normalize(),
        "Litros": d["LTS"].map(num),
        "CostoUnit": d["CTO UNIT"].map(num) if "CTO UNIT" in d.columns else 0.0,
        "Transportista": d.get("Transportista", ""),
        "Estatus": d.get("¿En repositorio?", ""),
    })
    return p.dropna(subset=["FechaBOL"]).reset_index(drop=True)

# ─────────────────────────── CONCILIACIÓN ───────────────────────────
def conciliar(pipas_ep, compras_ep):
    """Asigna pipas (fecha BOL, litros) a compras registradas en mermas.
    Una compra puede ser la suma de varias pipas (2 pipas descargadas el mismo día)."""
    libres = set(pipas_ep.index)
    asign = {}                       # idx pipa -> fecha registro
    compra_usada = {}                # idx compra -> [idx pipas]
    for w in range(VENTANA_EXTENDIDA + 1):
        for ci, c in compras_ep.iterrows():
            if ci in compra_usada: continue
            cands = [i for i in libres if abs((pipas_ep.at[i, "FechaBOL"] - c["Fecha"]).days) <= w]
            mejor = None
            for k in range(1, min(4, len(cands)) + 1):
                for combo in combinations(cands, k):
                    s = sum(pipas_ep.at[i, "Litros"] for i in combo)
                    if abs(s - c["Compras"]) <= TOLERANCIA_LTS * k:
                        score = sum(abs((pipas_ep.at[i, "FechaBOL"] - c["Fecha"]).days) for i in combo)
                        if mejor is None or score < mejor[0]:
                            mejor = (score, combo)
                if mejor: break
            if mejor:
                compra_usada[ci] = list(mejor[1])
                for i in mejor[1]:
                    asign[i] = c["Fecha"]; libres.discard(i)
    return asign, compra_usada

def procesar(origen):
    pipas_path, csvs = localizar_entradas(origen)
    if not pipas_path: raise FileNotFoundError("No encontré el Excel de PIPAS (el nombre debe contener 'PIPA').")
    pipas = leer_pipas(pipas_path)

    diario, avisos = [], []
    for f in sorted(csvs):
        nombre = os.path.splitext(os.path.basename(f))[0]
        est = estacion_oficial(nombre.split("-")[0])
        prod = producto_de(nombre.split("-", 1)[1] if "-" in nombre else nombre)
        if prod is None:
            avisos.append(f"No identifiqué el producto del archivo '{os.path.basename(f)}' (se omitió)."); continue
        if est not in ESTACIONES:
            avisos.append(f"La estación '{est}' del archivo '{os.path.basename(f)}' no está en la lista oficial (se incluyó igual).")
        if any((x.Estacion.iloc[0], x.Producto.iloc[0]) == (est, prod) for x in diario):
            avisos.append(f"Archivo duplicado para {est} {prod}: se ignoró '{f}'. Deja solo un CSV por estación y producto."); continue
        d, tot = leer_csv_mermas(f)
        d.insert(0, "Producto", prod); d.insert(0, "Estacion", est)
        if tot is not None and abs(d["MermaArchivo"].sum() - tot) > 1:
            avisos.append(f"{est} {prod}: la suma de mermas diarias ({d['MermaArchivo'].sum():,.2f}) no coincide con el TOTAL del archivo ({tot:,.2f}).")
        d["Nota"] = ""
        if len(d) and d.Fecha.iloc[0].day != 1:
            avisos.append(f"{est} {prod}: el reporte inicia el {d.Fecha.iloc[0]:%d/%m} (no trae los primeros días del mes). Las pipas anteriores a esa fecha se marcan FUERA DEL PERIODO.")
        if len(d) and abs(d.SdoInicial.iloc[0]) < 0.01 and d.SdoReal.iloc[0] > 0:
            i0 = d.index[0]
            if abs(d.at[i0, "Ajustes"]) < 0.01:
                ficticia = d.at[i0, "SdoReal"] - (d.at[i0, "Compras"] - d.at[i0, "Ventas"])
                d.at[i0, "SdoInicial"] = round(ficticia, 2)
                d.at[i0, "Nota"] = f"Sdo. Inicial venía en 0: se tomó el inventario real ({ficticia:,.0f} L) como saldo inicial"
                avisos.append(f"{est} {prod}: el {d.Fecha.iloc[0]:%d/%m} el Sdo. Inicial venía en 0, lo que generaba una 'merma' ficticia de {ficticia:+,.0f} L. "
                              f"Se tomó como saldo inicial el inventario real de ese día (Sdo. Real + Ventas − Compras). Corrige el arranque en el sistema.")
            else:
                d.at[i0, "Nota"] = "Sdo. Inicial en 0; el inventario de arranque viene en Ajustes"
                avisos.append(f"{est} {prod}: el {d.Fecha.iloc[0]:%d/%m} el Sdo. Inicial viene en 0 y el inventario de arranque se capturó como Ajuste ({d.at[i0, 'Ajustes']:,.0f} L).")
        diario.append(d)
    if not diario: raise FileNotFoundError("No encontré CSV de mermas.")
    diario = pd.concat(diario, ignore_index=True)
    # continuidad: Sdo.Inicial de hoy = Sdo.Real de ayer
    for (e, p), g in diario.groupby(["Estacion", "Producto"]):
        g = g.sort_values("Fecha")
        brk = (g["SdoInicial"].values[1:] - g["SdoReal"].values[:-1])
        for fch, b in zip(g["Fecha"].values[1:], brk):
            if abs(b) > 1:
                avisos.append(f"{e} {p}: el Sdo.Inicial del {pd.Timestamp(fch):%d/%m} no es igual al Sdo.Real del día anterior (dif {b:,.2f} L).")

    # Conciliación pipa por pipa
    det_pipas, compras_sin = [], []
    combos = set(zip(diario.Estacion, diario.Producto)) | set(zip(pipas.Estacion, pipas.Producto))
    fmin, fmax = diario.Fecha.min(), diario.Fecha.max()
    for (e, p) in sorted(combos):
        pe = pipas[(pipas.Estacion == e) & (pipas.Producto == p)].copy()
        de = diario[(diario.Estacion == e) & (diario.Producto == p)]
        if de.empty:
            for _, r in pe.iterrows():
                det_pipas.append({**r.to_dict(), "FechaRegistro": None, "Desfase": None, "Estado": "SIN REPORTE DE MERMAS"})
            continue
        fmin, fmax = de.Fecha.min(), de.Fecha.max()
        compras = de[de.Compras > 0][["Fecha", "Compras"]]
        asign, usadas = conciliar(pe, compras)
        # posibles errores de captura: compra sin pipa y pipa sin registro cercanas en fecha y litros
        captura = {}
        for ci, c in compras.iterrows():
            if ci in usadas: continue
            for i, r in pe.iterrows():
                if i in asign or i in captura.values(): continue
                if abs((c.Fecha - r.FechaBOL).days) <= VENTANA_DIAS and abs(c.Compras - r.Litros) <= UMBRAL_CAPTURA:
                    captura[ci] = i; break
        cap_p = {v: k for k, v in captura.items()}
        for i, r in pe.iterrows():
            if i in asign:
                dd = (asign[i] - r.FechaBOL).days
                est = "OK - mismo día" if dd == 0 else f"Registrada con desfase de {dd:+d} día(s)"
                if abs(dd) > VENTANA_DIAS: est += f" — FUERA DE TOLERANCIA ±{VENTANA_DIAS}"
                det_pipas.append({**r.to_dict(), "FechaRegistro": asign[i], "Desfase": dd, "Estado": est})
            else:
                nota = "NO REGISTRADA"
                if i in cap_p:
                    c = compras.loc[cap_p[i]]
                    nota = f"NO REGISTRADA con estos litros: posible error de captura (el {c.Fecha:%d/%m} se capturaron {c.Compras:,.0f} L; dif {c.Compras - r.Litros:+,.0f} L)"
                elif r.FechaBOL < fmin:
                    nota = f"FUERA DEL PERIODO DEL REPORTE (el reporte de mermas inicia el {fmin:%d/%m})"
                elif r.FechaBOL > fmax - pd.Timedelta(days=VENTANA_DIAS):
                    nota = "NO REGISTRADA (posible registro en el mes siguiente)"
                det_pipas.append({**r.to_dict(), "FechaRegistro": None, "Desfase": None, "Estado": nota})
        for ci, c in compras.iterrows():
            if ci not in usadas:
                nota = "Compra en mermas sin pipa/factura que la respalde"
                if ci in captura:
                    r = pe.loc[captura[ci]]
                    nota = f"Posible error de captura: la factura {r.Factura} ({r.FechaBOL:%d/%m}) es de {r.Litros:,.0f} L (dif {c.Compras - r.Litros:+,.0f} L)"
                elif c.Fecha < fmin + pd.Timedelta(days=VENTANA_DIAS):
                    nota += " (posible pipa del mes anterior)"
                compras_sin.append({"Estacion": e, "Producto": p, "Fecha": c.Fecha, "Litros": c.Compras, "Nota": nota})
    det_pipas = pd.DataFrame(det_pipas)
    compras_sin = pd.DataFrame(compras_sin, columns=["Estacion", "Producto", "Fecha", "Litros", "Nota"])
    return diario, pipas, det_pipas, compras_sin, avisos, sorted(combos), pipas_path, csvs

# ─────────────────────────── ESTILOS ───────────────────────────
AZUL = "1F3864"; AZUL2 = "D9E1F2"; GRIS = "F2F2F2"; MANUAL_F = "DDEBF7"; MANUAL_T = "1F4E79"
FILAS_MANUALES = 100      # renglones disponibles en la hoja "Registros manuales"
RM_LAST = 4 + FILAS_MANUALES
RM = {k: f"'Registros manuales'!${k}$5:${k}${RM_LAST}" for k in "ABCDEG"}
VERDE_F = "C6EFCE"; VERDE_T = "006100"; AMAR_F = "FFEB9C"; AMAR_T = "9C5700"; ROJO_F = "FFC7CE"; ROJO_T = "9C0006"
F = "Arial"
thin = Side(style="thin", color="BFBFBF")
BORDE = Border(left=thin, right=thin, top=thin, bottom=thin)
FMT_L = '#,##0.00;[Red]-#,##0.00;"-"'
FMT_L0 = '#,##0;[Red]-#,##0;"-"'
FMT_P = '0.00%;[Red]-0.00%;0.00%'
FMT_M = '$#,##0;[Red]-$#,##0;"-"'
FMT_D = "dd/mm/yyyy"

def titulo(ws, texto, sub, ancho):
    ws.sheet_view.showGridLines = False
    ws["A1"] = texto; ws["A1"].font = Font(name=F, size=16, bold=True, color="FFFFFF")
    ws["A2"] = sub;   ws["A2"].font = Font(name=F, size=10, italic=True, color="FFFFFF")
    for r in (1, 2):
        for c in range(1, ancho + 1):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=AZUL)
    ws.row_dimensions[1].height = 26

def encabezado(ws, fila, cols, anchos):
    for j, (c, w) in enumerate(zip(cols, anchos), 1):
        x = ws.cell(fila, j, c)
        x.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        x.fill = PatternFill("solid", fgColor="2F5597")
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        x.border = BORDE
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[fila].height = 32

def celda(ws, r, c, v, fmt=None, bold=False, align=None):
    x = ws.cell(r, c, v)
    x.font = Font(name=F, size=10, bold=bold); x.border = BORDE
    if fmt: x.number_format = fmt
    if align: x.alignment = Alignment(horizontal=align, vertical="center")
    return x

def semaforo(ws, rng, col_pct_abs_ref):
    """Relleno verde/amarillo/rojo según |Merma%| (col_pct_abs_ref = fórmula de la 1a celda)."""
    lim, ale = P_LIM, P_ALE
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({col_pct_abs_ref}),ABS({col_pct_abs_ref})>{lim})"],
        fill=PatternFill("solid", fgColor=ROJO_F), font=Font(color=ROJO_T, bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({col_pct_abs_ref}),ABS({col_pct_abs_ref})>{ale})"],
        fill=PatternFill("solid", fgColor=AMAR_F), font=Font(color=AMAR_T, bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"ISNUMBER({col_pct_abs_ref})"],
        fill=PatternFill("solid", fgColor=VERDE_F), font=Font(color=VERDE_T, bold=True), stopIfTrue=True))

def estado_rule(ws, rng, first):
    for txt, fi, ft in (("FUERA", ROJO_F, ROJO_T), ("ALERTA", AMAR_F, AMAR_T), ("OK", VERDE_F, VERDE_T)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="{txt}"'],
            fill=PatternFill("solid", fgColor=fi), font=Font(color=ft, bold=True)))

def texto_rule(ws, rng, first, palabra, fi, ft):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{palabra}",{first}))'],
        fill=PatternFill("solid", fgColor=fi), font=Font(color=ft, bold=True)))

def foco(ws, rng):
    ws.conditional_formatting.add(rng, IconSetRule("3TrafficLights1", "num", [0, UMBRAL_ALERTA, LIMITE_SAT],
                                                   showValue=False, reverse=True))

# ─────────────────────────── EXCEL ───────────────────────────
def construir_excel(salida, diario, pipas, det_pipas, compras_sin, avisos, combos, pipas_path, csvs):
    wb = Workbook()
    mes = diario.Fecha.min()
    MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
             "Septiembre", "Octubre", "Noviembre", "Diciembre"]
    periodo = f"{MESES[mes.month]} {mes.year}"
    gen = f"Periodo: {periodo}  ·  Generado: {datetime.now():%d/%m/%Y %H:%M}  ·  Límite SAT: merma ≤ 0.5% de la venta"

    ws_res = wb.active; ws_res.title = "Resumen"
    ws_mm = wb.create_sheet("Mermas mensual")
    ws_s = wb.create_sheet("Compras sin pipa")
    ws_d = wb.create_sheet("Detalle diario")
    ws_c = wb.create_sheet("Conciliación pipas")
    ws_p = wb.create_sheet("Detalle pipas")
    ws_m = wb.create_sheet("Registros manuales")
    ws_par = wb.create_sheet("Parámetros")
    ws_n = wb.create_sheet("Notas")

    # ---------- Parámetros ----------
    titulo(ws_par, "Parámetros", "Cambia los valores en amarillo y todo el libro se recalcula", 4)
    encabezado(ws_par, 4, ["", "Parámetro", "Valor", "Descripción"], [2, 34, 12, 70])
    params = [("Límite SAT merma / venta", LIMITE_SAT, FMT_P, "Arriba de este % el foco es ROJO (FUERA)"),
              ("Umbral de alerta", UMBRAL_ALERTA, FMT_P, "Entre este % y el límite el foco es AMARILLO (ALERTA)"),
              ("Ventana de desfase (días)", VENTANA_DIAS, "0", "Días ± entre fecha BOL y registro en mermas (se usó al generar)"),
              ("Tolerancia por pipa (L)", TOLERANCIA_LTS, "0", "Diferencia máxima en litros para considerar que una pipa cuadra"),
              ("Umbral salto diario (L)", UMBRAL_SALTO, "#,##0", "|Merma diaria| mayor a esto = posible pipa registrada en otro día")]
    for i, (a, b, fmt, c) in enumerate(params, 5):
        celda(ws_par, i, 2, a, bold=True); x = celda(ws_par, i, 3, b, fmt, align="center")
        x.fill = PatternFill("solid", fgColor="FFFF00"); x.font = Font(name=F, size=10, color="0000FF", bold=True)
        celda(ws_par, i, 4, c)

    # ---------- Costo promedio por estación/producto ----------
    pp = pipas.assign(Imp=pipas.Litros * pipas.CostoUnit).groupby(["Estacion", "Producto"]).agg(L=("Litros", "sum"), I=("Imp", "sum"), N=("Litros", "size"))
    pp["Costo"] = pp.I / pp.L

    # ---------- Detalle diario ----------
    titulo(ws_d, "Detalle diario de mermas", gen, 12)
    cols = ["Estación", "Producto", "Fecha", "Sdo. Inicial", "Compras", "Ventas", "Ajustes", "Sdo. Final",
            "Sdo. Real", "Merma", "Merma %", "Compras\nmanuales"]
    encabezado(ws_d, 4, cols, [14, 10, 11, 13, 12, 12, 9, 13, 13, 11, 10, 12])
    r = 5
    dsort = diario.assign(o=diario.Estacion.map(lambda x: ESTACIONES.index(x) if x in ESTACIONES else 99),
                          q=diario.Producto.map(PRODUCTOS.index)).sort_values(["o", "Estacion", "q", "Fecha"])
    for _, x in dsort.iterrows():
        celda(ws_d, r, 1, x.Estacion); celda(ws_d, r, 2, x.Producto); celda(ws_d, r, 3, x.Fecha.to_pydatetime(), FMT_D, align="center")
        for j, k in ((4, "SdoInicial"), (5, "Compras"), (6, "Ventas"), (7, "Ajustes"), (9, "SdoReal")):
            celda(ws_d, r, j, round(x[k], 2), FMT_L)
        celda(ws_d, r, 8, f"=D{r}+E{r}+L{r}-F{r}+G{r}", FMT_L)
        celda(ws_d, r, 12, f"=SUMIFS({RM['D']},{RM['A']},A{r},{RM['B']},B{r},{RM['C']},C{r})", FMT_L)
        celda(ws_d, r, 10, f"=I{r}-H{r}", FMT_L)
        celda(ws_d, r, 11, f"=IF(F{r}=0,0,J{r}/F{r})", FMT_P)
        r += 1
    last_d = r - 1
    semaforo(ws_d, f"K5:K{last_d}", "K5")
    ws_d.conditional_formatting.add(f"L5:L{last_d}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=MANUAL_F), font=Font(color=MANUAL_T, bold=True)))
    ws_d.freeze_panes = "D5"; ws_d.auto_filter.ref = f"A4:L{last_d}"

    # ---------- Mermas mensual ----------
    # A Est · B Prod · C SI · D Compras · E Ventas · F Ajustes · G SF · H SR · I Merma · J Merma% · K Estado (lo usa Resumen)
    # L Costo prom (lo usa Merma $) · M Merma $ (lo usa Resumen) · N Pipas no reg · O Merma ajustada · P Merma % ajustada
    titulo(ws_mm, "Mermas mensuales por estación y producto", gen, 16)
    cols = ["Estación", "Producto", "Sdo. Inicial\n(1er día)", "Compras", "Ventas", "Ajustes", "Sdo. Final\n(teórico)",
            "Sdo. Real\n(último día)", "Merma (L)", "Merma %", "Estado", "Costo prom.\npipas ($/L)", "Merma ($)",
            "Pipas NO\nregistradas (L)", "Merma ajustada\n(L)", "Merma %\najustada"]
    encabezado(ws_mm, 4, cols, [14, 10, 13, 13, 13, 10, 13, 13, 12, 10, 9, 11, 12, 13, 13, 10])
    r = 5
    agr = []
    nP_ = len(det_pipas) + 4
    for e in ESTACIONES + sorted({c[0] for c in combos} - set(ESTACIONES)):
        for p in PRODUCTOS:
            g = diario[(diario.Estacion == e) & (diario.Producto == p)].sort_values("Fecha")
            if g.empty: continue
            celda(ws_mm, r, 1, e, bold=True); celda(ws_mm, r, 2, p)
            celda(ws_mm, r, 3, round(g.SdoInicial.iloc[0], 2), FMT_L)
            rngA, rngB = f"'Detalle diario'!$A$5:$A${last_d}", f"'Detalle diario'!$B$5:$B${last_d}"
            celda(ws_mm, r, 4, f"=SUMIFS('Detalle diario'!$E$5:$E${last_d},{rngA},A{r},{rngB},B{r})+SUMIFS('Detalle diario'!$L$5:$L${last_d},{rngA},A{r},{rngB},B{r})", FMT_L)
            celda(ws_mm, r, 5, f"=SUMIFS('Detalle diario'!$F$5:$F${last_d},{rngA},A{r},{rngB},B{r})", FMT_L)
            celda(ws_mm, r, 6, f"=SUMIFS('Detalle diario'!$G$5:$G${last_d},{rngA},A{r},{rngB},B{r})", FMT_L)
            celda(ws_mm, r, 7, f"=C{r}+D{r}-E{r}+F{r}", FMT_L)
            celda(ws_mm, r, 8, round(g.SdoReal.iloc[-1], 2), FMT_L)
            celda(ws_mm, r, 9, f"=H{r}-G{r}", FMT_L, bold=True)
            celda(ws_mm, r, 10, f"=IF(E{r}=0,0,I{r}/E{r})", FMT_P, bold=True)
            celda(ws_mm, r, 11, f'=IF(ABS(J{r})>{P_LIM},"FUERA",IF(ABS(J{r})>{P_ALE},"ALERTA","OK"))', align="center")
            cst = pp.Costo.get((e, p), None)
            x = celda(ws_mm, r, 12, round(float(cst), 4) if cst is not None else 0, '$#,##0.00')
            x.font = Font(name=F, size=10, color="0000FF")
            celda(ws_mm, r, 13, f"=I{r}*L{r}", FMT_M)
            celda(ws_mm, r, 14, f"=SUMIFS('Detalle pipas'!$G$5:$G${nP_},'Detalle pipas'!$A$5:$A${nP_},A{r},'Detalle pipas'!$B$5:$B${nP_},B{r},'Detalle pipas'!$L$5:$L${nP_},\"NO REGISTRADA*\")", FMT_L0)
            celda(ws_mm, r, 15, f"=I{r}-N{r}", FMT_L)
            celda(ws_mm, r, 16, f"=IF(E{r}=0,0,O{r}/E{r})", FMT_P)
            agr.append((e, p, r)); r += 1
    last_mm = r - 1
    # Total
    celda(ws_mm, r, 1, "TOTAL GRUPO", bold=True); celda(ws_mm, r, 2, "")
    for j in (3, 4, 5, 6, 7, 8, 9, 13, 14, 15):
        L = get_column_letter(j); celda(ws_mm, r, j, f"=SUM({L}5:{L}{last_mm})", FMT_M if j == 13 else FMT_L, bold=True)
    celda(ws_mm, r, 10, f"=IF(E{r}=0,0,I{r}/E{r})", FMT_P, bold=True)
    celda(ws_mm, r, 11, f'=IF(ABS(J{r})>{P_LIM},"FUERA",IF(ABS(J{r})>{P_ALE},"ALERTA","OK"))', align="center")
    celda(ws_mm, r, 12, "")
    celda(ws_mm, r, 16, f"=IF(E{r}=0,0,O{r}/E{r})", FMT_P, bold=True)
    for j in range(1, 17): ws_mm.cell(r, j).fill = PatternFill("solid", fgColor=AZUL2)
    tot_mm = r
    estado_rule(ws_mm, f"K5:K{tot_mm}", "K5")
    semaforo(ws_mm, f"J5:J{tot_mm}", "J5")
    semaforo(ws_mm, f"P5:P{tot_mm}", "P5")
    ws_mm.conditional_formatting.add(f"N5:N{last_mm}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=ROJO_F), font=Font(color=ROJO_T, bold=True)))
    ws_mm.freeze_panes = "C5"; ws_mm.auto_filter.ref = f"A4:P{last_mm}"
    ws_mm.cell(tot_mm + 5, 1, "Compras incluye lo capturado en la hoja Registros manuales.").font = Font(name=F, size=9, italic=True)
    ws_mm.cell(tot_mm + 2, 1, "Merma = Sdo. Real − Sdo. Final;  Sdo. Final = Sdo. Inicial + Compras − Ventas + Ajustes;  Merma % = Merma / Ventas.  "
               "Negativo = faltante, positivo = sobrante.").font = Font(name=F, size=9, italic=True)
    ws_mm.cell(tot_mm + 3, 1, "Costo prom. = promedio ponderado del CTO UNIT de las pipas del mes (dato del Excel de pipas); Merma ($) = Merma (L) × Costo prom.").font = Font(name=F, size=9, italic=True)
    ws_mm.cell(tot_mm + 4, 1, "Merma ajustada = Merma − litros de pipas NO registradas: si una pipa llegó físicamente pero no se capturó, aparece como sobrante. Es un indicador de apoyo; el oficial es Merma %.").font = Font(name=F, size=9, italic=True)


    # ---------- Conciliación mensual ----------
    titulo(ws_c, "Conciliación de compras: Reporte de mermas vs PIPAS DEL MES", gen, 11)
    cols = ["Estación", "Producto", "Compras\nreporte mermas (L)", "Litros\nPIPAS DEL MES", "# Pipas", "Diferencia (L)\nmermas − pipas",
            "Estado", "Pipas no\nregistradas", "Pipas con\ndesfase", "Compras\nsin pipa", "Comentario"]
    encabezado(ws_c, 4, cols, [14, 10, 16, 15, 8, 15, 16, 11, 11, 10, 80])
    r = 5
    orden = sorted(combos, key=lambda c: (ESTACIONES.index(c[0]) if c[0] in ESTACIONES else 99, c[0], PRODUCTOS.index(c[1]) if c[1] in PRODUCTOS else 9))
    mm_row = {(e, p): rr for e, p, rr in agr}
    nP = len(det_pipas) + 4; nS = max(len(compras_sin) + 4, 5)
    for e, p in orden:
        celda(ws_c, r, 1, e, bold=True); celda(ws_c, r, 2, p)
        if (e, p) in mm_row:
            celda(ws_c, r, 3, f"='Mermas mensual'!D{mm_row[(e, p)]}", FMT_L).font = Font(name=F, size=10, color="008000")
        else:
            celda(ws_c, r, 3, "Sin reporte", align="center")
        celda(ws_c, r, 4, f"=SUMIFS('Detalle pipas'!$G$5:$G${nP},'Detalle pipas'!$A$5:$A${nP},A{r},'Detalle pipas'!$B$5:$B${nP},B{r},'Detalle pipas'!$M$5:$M${nP},\"Sí\")", FMT_L0)
        celda(ws_c, r, 5, f"=COUNTIFS('Detalle pipas'!$A$5:$A${nP},A{r},'Detalle pipas'!$B$5:$B${nP},B{r},'Detalle pipas'!$M$5:$M${nP},\"Sí\")", "0", align="center")
        celda(ws_c, r, 6, f'=IF(ISNUMBER(C{r}),C{r}-D{r},"")', FMT_L, bold=True)
        celda(ws_c, r, 7, f'=IF(NOT(ISNUMBER(C{r})),"SIN REPORTE",IF(ABS(F{r})<={P_TOL}*MAX(1,E{r}),"CUADRA","NO CUADRA"))', align="center")
        celda(ws_c, r, 8, f"=COUNTIFS('Detalle pipas'!$A$5:$A${nP},A{r},'Detalle pipas'!$B$5:$B${nP},B{r},'Detalle pipas'!$L$5:$L${nP},\"NO REGISTRADA*\")", "0", align="center")
        celda(ws_c, r, 9, f"=COUNTIFS('Detalle pipas'!$A$5:$A${nP},A{r},'Detalle pipas'!$B$5:$B${nP},B{r},'Detalle pipas'!$L$5:$L${nP},\"Registrada con desfase*\")", "0", align="center")
        celda(ws_c, r, 10, f"=COUNTIFS('Compras sin pipa'!$A$5:$A${nS},A{r},'Compras sin pipa'!$B$5:$B${nS},B{r})"
                             f"+COUNTIFS({RM['A']},A{r},{RM['B']},B{r},{RM['G']},\"SIN PIPA*\")", "0", align="center")
        celda(ws_c, r, 11, f'=IF(COUNTIFS(\'Detalle pipas\'!$A$5:$A${nP},A{r},\'Detalle pipas\'!$B$5:$B${nP},B{r},\'Detalle pipas\'!$L$5:$L${nP},"REGISTRO MANUAL*")>0,'
                             f'"Incluye "&COUNTIFS(\'Detalle pipas\'!$A$5:$A${nP},A{r},\'Detalle pipas\'!$B$5:$B${nP},B{r},\'Detalle pipas\'!$L$5:$L${nP},"REGISTRO MANUAL*")&" registro(s) manual(es) · ","")&IF(G{r}="SIN REPORTE","Falta el CSV de mermas de esta estación/producto",'
                             f'IF(G{r}="CUADRA",IF(I{r}>0,"Cuadra en el mes; hay pipas registradas en otro día (ver Detalle pipas)","Todo en orden"),'
                             f'"Revisar: "&H{r}&" pipa(s) sin registrar y "&J{r}&" compra(s) sin pipa"))'
                             f'&IF(COUNTIFS(\'Detalle pipas\'!$A$5:$A${nP},A{r},\'Detalle pipas\'!$B$5:$B${nP},B{r},\'Detalle pipas\'!$M$5:$M${nP},"No")>0,'
                             f'" · "&COUNTIFS(\'Detalle pipas\'!$A$5:$A${nP},A{r},\'Detalle pipas\'!$B$5:$B${nP},B{r},\'Detalle pipas\'!$M$5:$M${nP},"No")&" pipa(s) antes del inicio del reporte (no se cuentan)","")')
        r += 1
    last_c = r - 1
    for txt, fi, ft in (("NO CUADRA", ROJO_F, ROJO_T), ("SIN REPORTE", GRIS, "595959"), ("CUADRA", VERDE_F, VERDE_T)):
        ws_c.conditional_formatting.add(f"G5:G{last_c}", FormulaRule(formula=[f'G5="{txt}"'], fill=PatternFill("solid", fgColor=fi), font=Font(color=ft, bold=True)))
    for L in ("H", "J"):
        ws_c.conditional_formatting.add(f"{L}5:{L}{last_c}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=ROJO_F), font=Font(color=ROJO_T, bold=True)))
    ws_c.conditional_formatting.add(f"I5:I{last_c}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=AMAR_F), font=Font(color=AMAR_T, bold=True)))
    ws_c.freeze_panes = "C5"; ws_c.auto_filter.ref = f"A4:K{last_c}"

    # ---------- Detalle pipas ----------
    titulo(ws_p, "Detalle pipa por pipa (PIPAS DEL MES vs registro en mermas)", gen, 14)
    cols = ["Estación", "Producto", "Factura", "BOL", "Fecha BOL", "Transportista", "Litros", "Costo unit.", "Fecha registro\nen mermas",
            "Desfase\n(días)", "Importe ($)", "Estado", "En periodo\ndel reporte", "Estado original\n(sin registros manuales)"]
    encabezado(ws_p, 4, cols, [14, 10, 10, 10, 11, 13, 11, 10, 13, 9, 13, 60, 11, 40])
    r = 5
    if len(det_pipas):
        dps = det_pipas.assign(o=det_pipas.Estacion.map(lambda x: ESTACIONES.index(x) if x in ESTACIONES else 99)).sort_values(["o", "Estacion", "Producto", "FechaBOL"])
        for x in dps.itertuples():
            celda(ws_p, r, 1, x.Estacion); celda(ws_p, r, 2, x.Producto)
            celda(ws_p, r, 3, int(x.Factura) if pd.notna(x.Factura) else "", align="center")
            celda(ws_p, r, 4, int(x.BOL) if pd.notna(x.BOL) else "", align="center")
            celda(ws_p, r, 5, x.FechaBOL.to_pydatetime(), FMT_D, align="center")
            celda(ws_p, r, 6, x.Transportista if pd.notna(x.Transportista) else "")
            celda(ws_p, r, 7, x.Litros, FMT_L0); celda(ws_p, r, 8, x.CostoUnit, '$#,##0.0000')
            celda(ws_p, r, 11, f"=G{r}*H{r}", FMT_M)
            celda(ws_p, r, 14, x.Estado).font = Font(name=F, size=9, color="7F7F7F")
            if str(x.Estado).startswith("NO REGISTRADA"):
                # se vuelve REGISTRO MANUAL si aparece en la hoja Registros manuales (por factura, o por litros ± tolerancia sin factura)
                mf = f"COUNTIFS({RM['E']},C{r},{RM['A']},A{r},{RM['B']},B{r})"
                ml = f'COUNTIFS({RM["A"]},A{r},{RM["B"]},B{r},{RM["E"]},"",{RM["D"]},">="&(G{r}-{P_TOL}),{RM["D"]},"<="&(G{r}+{P_TOL}))'
                fm = f"SUMIFS({RM['C']},{RM['E']},C{r},{RM['A']},A{r},{RM['B']},B{r})"
                fl = f'SUMIFS({RM["C"]},{RM["A"]},A{r},{RM["B"]},B{r},{RM["E"]},"",{RM["D"]},">="&(G{r}-{P_TOL}),{RM["D"]},"<="&(G{r}+{P_TOL}))'
                celda(ws_p, r, 9, f'=IF({mf}>0,{fm},IF({ml}>0,{fl},""))', FMT_D, align="center")
                celda(ws_p, r, 10, f'=IF(I{r}="","",I{r}-E{r})', '+0;-0;0', align="center")
                celda(ws_p, r, 12, f'=IF(OR({mf}>0,{ml}>0),"REGISTRO MANUAL el "&TEXT(I{r},"dd/mm"),N{r})')
            else:
                celda(ws_p, r, 9, x.FechaRegistro.to_pydatetime() if pd.notna(x.FechaRegistro) else "", FMT_D, align="center")
                celda(ws_p, r, 10, int(x.Desfase) if pd.notna(x.Desfase) else "", '+0;-0;0', align="center")
                celda(ws_p, r, 12, x.Estado)
            celda(ws_p, r, 13, "No" if str(x.Estado).startswith("FUERA DEL PERIODO") else "Sí", align="center")
            r += 1
    last_p = max(r - 1, 5)
    texto_rule(ws_p, f"L5:L{last_p}", "L5", "REGISTRO MANUAL", MANUAL_F, MANUAL_T)
    texto_rule(ws_p, f"L5:L{last_p}", "L5", "FUERA DE TOLERANCIA", "F8CBAD", "833C0B")
    texto_rule(ws_p, f"L5:L{last_p}", "L5", "FUERA DEL PERIODO", GRIS, "595959")
    texto_rule(ws_p, f"L5:L{last_p}", "L5", "NO REGISTRADA", ROJO_F, ROJO_T)
    texto_rule(ws_p, f"L5:L{last_p}", "L5", "desfase", AMAR_F, AMAR_T)
    texto_rule(ws_p, f"L5:L{last_p}", "L5", "SIN REPORTE", GRIS, "595959")
    texto_rule(ws_p, f"L5:L{last_p}", "L5", "OK", VERDE_F, VERDE_T)
    ws_p.freeze_panes = "C5"; ws_p.auto_filter.ref = f"A4:N{last_p}"

    # ---------- Compras sin pipa ----------
    titulo(ws_s, "Compras registradas en mermas que no se encontraron en PIPAS DEL MES", gen, 5)
    encabezado(ws_s, 4, ["Estación", "Producto", "Fecha", "Litros", "Nota"], [14, 10, 12, 12, 60])
    r = 5
    for x in compras_sin.itertuples():
        celda(ws_s, r, 1, x.Estacion); celda(ws_s, r, 2, x.Producto); celda(ws_s, r, 3, x.Fecha.to_pydatetime(), FMT_D, align="center")
        celda(ws_s, r, 4, x.Litros, FMT_L); celda(ws_s, r, 5, x.Nota); r += 1
    if r == 5:
        ws_s.cell(5, 1, "✔ Todas las compras del reporte de mermas tienen una pipa que las respalda.").font = Font(name=F, color=VERDE_T, bold=True)
    ws_s.freeze_panes = "A5"

    # ---------- Registros manuales (captura dentro del reporte) ----------
    titulo(ws_m, "Registros manuales de pipas", "Captura aquí las pipas que faltaban por registrar (celdas amarillas). Todo el libro se recalcula solo; no hay que volver a correr el script.", 7)
    encabezado(ws_m, 4, ["Estación", "Producto", "Fecha registro", "Litros", "Factura\n(opcional)", "Comentario", "Resultado"], [16, 12, 14, 12, 12, 40, 58])
    DA, DB, DC = (f"'Detalle diario'!${k}$5:${k}${last_d}" for k in "ABC")
    PA, PB, PC, PG, PN = (f"'Detalle pipas'!${k}$5:${k}${last_p}" for k in "ABCGN")
    for r in range(5, RM_LAST + 1):
        for j in range(1, 7):
            x = ws_m.cell(r, j); x.border = BORDE; x.font = Font(name=F, size=10, color="0000FF")
            x.fill = PatternFill("solid", fgColor="FFFF00" if j <= 4 else "FFFFCC")
        ws_m.cell(r, 3).number_format = FMT_D; ws_m.cell(r, 4).number_format = "#,##0"
        nfr = f"COUNTIFS({PC},E{r},{PA},A{r},{PB},B{r},{PN},\"NO REGISTRADA*\")"
        nf = f"COUNTIFS({PC},E{r},{PA},A{r},{PB},B{r})"
        nl = f'COUNTIFS({PA},A{r},{PB},B{r},{PG},">="&(D{r}-{P_TOL}),{PG},"<="&(D{r}+{P_TOL}),{PN},"NO REGISTRADA*")'
        nd = f"COUNTIFS({DA},A{r},{DB},B{r},{DC},C{r})"
        celda(ws_m, r, 7, f'=IF(AND(A{r}="",D{r}=""),"",IF(OR(A{r}="",B{r}="",C{r}="",D{r}=""),"FALTAN DATOS: Estación, Producto, Fecha y Litros son obligatorios",'
                          f'IF({nd}=0,"NO APLICA: no hay reporte de mermas de esa estación/producto en esa fecha",'
                          f'IF(E{r}<>"",IF({nfr}>0,"APLICADO: registra la factura "&E{r},IF({nf}>0,"OJO: la factura "&E{r}&" ya estaba registrada; se estaría sumando doble","SIN PIPA: la factura "&E{r}&" no existe en PIPAS DEL MES para esa estación/producto")),'
                          f'IF({nl}>0,"APLICADO: registra la pipa de "&TEXT(D{r},"#,##0")&" L","SIN PIPA: no hay pipa pendiente con esos litros (revisa litros o pon la factura)")))))')
    dv = DataValidation(type="list", formula1='"' + ",".join(ESTACIONES) + '"', allow_blank=True); dv.add(f"A5:A{RM_LAST}")
    dv2 = DataValidation(type="list", formula1='"Regular,Premium,Diesel"', allow_blank=True); dv2.add(f"B5:B{RM_LAST}")
    dv3 = DataValidation(type="date", operator="greaterThan", formula1="40000", allow_blank=True, showErrorMessage=True, error="Escribe una fecha dd/mm/aaaa"); dv3.add(f"C5:C{RM_LAST}")
    dv4 = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True, showErrorMessage=True, error="Litros debe ser mayor a 0"); dv4.add(f"D5:D{RM_LAST}")
    for v in (dv, dv2, dv3, dv4): ws_m.add_data_validation(v)
    rg = f"G5:G{RM_LAST}"
    texto_rule(ws_m, rg, "G5", "APLICADO", VERDE_F, VERDE_T)
    texto_rule(ws_m, rg, "G5", "OJO", AMAR_F, AMAR_T)
    for pal in ("SIN PIPA", "FALTAN", "NO APLICA"):
        texto_rule(ws_m, rg, "G5", pal, ROJO_F, ROJO_T)
    ws_m.freeze_panes = "A5"
    nota_m = ws_m.cell(RM_LAST + 2, 1, "Fecha registro = día del reporte de mermas en que se suma la compra. Con Factura, se liga a esa pipa; sin Factura, se busca una pipa NO REGISTRADA "
                                        "de la misma estación/producto con esos litros (± tolerancia de Parámetros).")
    nota_m.font = Font(name=F, size=9, italic=True)

    # ---------- Resumen ----------
    W = 9
    titulo(ws_res, f"Tablero de mermas · {periodo}", gen, W)
    ws_res.column_dimensions["A"].width = 16
    for c in "BCDEFGHI": ws_res.column_dimensions[c].width = 15
    # KPIs
    kp = [("Merma total grupo (L)", f"='Mermas mensual'!I{tot_mm}", FMT_L),
          ("Merma total grupo %", f"='Mermas mensual'!J{tot_mm}", FMT_P),
          ("Merma total grupo ($)", f"='Mermas mensual'!M{tot_mm}", FMT_M),
          ("Combinaciones FUERA", f"=COUNTIF('Mermas mensual'!K5:K{last_mm},\"FUERA\")", "0"),
          ("Pipas NO registradas", f"=COUNTIF('Detalle pipas'!L5:L{last_p},\"NO REGISTRADA*\")", "0"),
          ("Compras sin pipa", f"=COUNTA('Compras sin pipa'!D5:D{max(len(compras_sin)+4,5)})+COUNTIF({RM['G']},\"SIN PIPA*\")", "0")]
    for i, (lab, fml, fmt) in enumerate(kp):
        col = 1 + i * (W // 6 if i < 3 else 1)
    # simpler fixed layout: 2 rows x 3 KPIs, each spans 3 columns
    pos = [(4, 1), (4, 4), (4, 7), (7, 1), (7, 4), (7, 7)]
    for (lab, fml, fmt), (rr, cc) in zip(kp, pos):
        ws_res.merge_cells(start_row=rr, start_column=cc, end_row=rr, end_column=cc + 2)
        ws_res.merge_cells(start_row=rr + 1, start_column=cc, end_row=rr + 1, end_column=cc + 2)
        a = ws_res.cell(rr, cc, lab); a.font = Font(name=F, size=9, bold=True, color="595959"); a.alignment = Alignment(horizontal="center")
        b = ws_res.cell(rr + 1, cc, fml); b.font = Font(name=F, size=18, bold=True, color=AZUL); b.number_format = fmt
        b.alignment = Alignment(horizontal="center", vertical="center")
        for k in range(3):
            ws_res.cell(rr, cc + k).fill = PatternFill("solid", fgColor=GRIS)
            ws_res.cell(rr + 1, cc + k).fill = PatternFill("solid", fgColor=GRIS)
        ws_res.row_dimensions[rr + 1].height = 30
    ws_res.conditional_formatting.add("D5", FormulaRule(formula=[f"ABS(D5)>{P_LIM}"], font=Font(color=ROJO_T, bold=True, size=18)))
    for ref in ("D8", "G8", "J8"):
        pass
    for ref in ("A8", "D8", "G8"):
        ws_res.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=ROJO_T, bold=True, size=18)))

    # Matriz Merma %
    r0 = 11
    ws_res.cell(r0 - 1, 1, "Merma % mensual por estación y producto").font = Font(name=F, size=12, bold=True, color=AZUL)
    hdr = ["Estación", "Regular", "Premium", "Diesel", "Total estación", "Merma (L)", "Merma ($)", "Compras vs pipas", "Pipas sin registrar"]
    for j, h in enumerate(hdr, 1):
        x = ws_res.cell(r0, j, h); x.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        x.fill = PatternFill("solid", fgColor="2F5597"); x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); x.border = BORDE
    ws_res.row_dimensions[r0].height = 30
    MA, MB = f"'Mermas mensual'!$A$5:$A${last_mm}", f"'Mermas mensual'!$B$5:$B${last_mm}"
    MI, ME, MP = f"'Mermas mensual'!$I$5:$I${last_mm}", f"'Mermas mensual'!$E$5:$E${last_mm}", f"'Mermas mensual'!$M$5:$M${last_mm}"
    CA, CG = f"'Conciliación pipas'!$A$5:$A${last_c}", f"'Conciliación pipas'!$G$5:$G${last_c}"
    CH = f"'Conciliación pipas'!$H$5:$H${last_c}"
    r = r0 + 1
    for e in ESTACIONES:
        celda(ws_res, r, 1, e, bold=True)
        for j, p in enumerate(PRODUCTOS, 2):
            celda(ws_res, r, j, f'=IF(COUNTIFS({MA},$A{r},{MB},"{p}")=0,"s/d",IF(SUMIFS({ME},{MA},$A{r},{MB},"{p}")=0,0,'
                                f'SUMIFS({MI},{MA},$A{r},{MB},"{p}")/SUMIFS({ME},{MA},$A{r},{MB},"{p}")))', FMT_P, align="center")
        celda(ws_res, r, 5, f'=IF(COUNTIF({MA},$A{r})=0,"s/d",IF(SUMIFS({ME},{MA},$A{r})=0,0,SUMIFS({MI},{MA},$A{r})/SUMIFS({ME},{MA},$A{r})))', FMT_P, bold=True, align="center")
        celda(ws_res, r, 6, f'=IF(COUNTIF({MA},$A{r})=0,"s/d",SUMIFS({MI},{MA},$A{r}))', FMT_L, align="right")
        celda(ws_res, r, 7, f'=IF(COUNTIF({MA},$A{r})=0,"s/d",SUMIFS({MP},{MA},$A{r}))', FMT_M, align="right")
        celda(ws_res, r, 8, f'=IF(COUNTIFS({CA},$A{r},{CG},"NO CUADRA")>0,"NO CUADRA",IF(COUNTIFS({CA},$A{r},{CG},"CUADRA")>0,'
                            f'IF(COUNTIFS({CA},$A{r},{CG},"SIN REPORTE")>0,"PARCIAL","CUADRA"),IF(COUNTIF({CA},$A{r})=0,"s/d","SIN REPORTE")))', align="center")
        celda(ws_res, r, 9, f'=SUMIFS({CH},{CA},$A{r})', "0", align="center")
        r += 1
    last_r = r - 1
    semaforo(ws_res, f"B{r0+1}:E{last_r}", f"B{r0+1}")
    ws_res.conditional_formatting.add(f"B{r0+1}:G{last_r}", FormulaRule(formula=[f'B{r0+1}="s/d"'], font=Font(color="A6A6A6", italic=True)))
    for txt, fi, ft in (("NO CUADRA", ROJO_F, ROJO_T), ("PARCIAL", AMAR_F, AMAR_T), ("SIN REPORTE", GRIS, "595959"), ("CUADRA", VERDE_F, VERDE_T)):
        ws_res.conditional_formatting.add(f"H{r0+1}:H{last_r}", FormulaRule(formula=[f'H{r0+1}="{txt}"'], fill=PatternFill("solid", fgColor=fi), font=Font(color=ft, bold=True)))
    ws_res.conditional_formatting.add(f"I{r0+1}:I{last_r}", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=ROJO_F), font=Font(color=ROJO_T, bold=True)))
    lg = last_r + 2
    ws_res.cell(lg, 1, "Semáforo:").font = Font(name=F, bold=True, size=9)
    for j, (t, fi, ft) in enumerate((("Verde: |merma| ≤ 0.40 %", VERDE_F, VERDE_T), ("Amarillo: 0.40 % – 0.50 %", AMAR_F, AMAR_T),
                                     ("Rojo: > 0.50 % (fuera SAT)", ROJO_F, ROJO_T), ("s/d: sin reporte de mermas", GRIS, "595959"))):
        ws_res.merge_cells(start_row=lg, start_column=2 + j * 2, end_row=lg, end_column=3 + j * 2)
        x = ws_res.cell(lg, 2 + j * 2, t); x.fill = PatternFill("solid", fgColor=fi); x.font = Font(name=F, size=9, bold=True, color=ft)
        x.alignment = Alignment(horizontal="center")
    ws_res.cell(lg + 1, 1, "Negativo = faltante de combustible · Positivo = sobrante. Los umbrales se cambian en la hoja Parámetros.").font = Font(name=F, size=9, italic=True)

    # Gráfica
    ch = BarChart(); ch.type = "bar"; ch.style = 10
    ch.title = "Merma % por estación y producto"; ch.y_axis.title = "Merma % de la venta"
    ch.y_axis.numFmt = "0.00%"; ch.x_axis.tickLblPos = "low"; ch.x_axis.delete = False; ch.y_axis.delete = False; ch.height = 7 + 0.45 * len(agr); ch.width = 18; ch.legend = None
    data = Reference(ws_mm, min_col=10, min_row=4, max_row=last_mm)
    cats_ws = ws_mm
    # etiqueta combinada Estación-Producto en columna oculta R
    for rr in range(5, last_mm + 1):
        ws_mm.cell(rr, 18, f'=A{rr}&" · "&B{rr}')
        ws_mm.cell(rr, 18).font = Font(name=F, size=8, color="A6A6A6")
    ws_mm.cell(4, 18, "Etiqueta\ngráfica").font = Font(name=F, size=8, color="A6A6A6")
    ws_mm.column_dimensions["R"].width = 22
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(Reference(ws_mm, min_col=18, min_row=5, max_row=last_mm))
    ws_res.add_chart(ch, f"A{lg + 3}")

    # ---------- Notas ----------
    titulo(ws_n, "Notas, metodología y avisos", gen, 2)
    ws_n.column_dimensions["A"].width = 4; ws_n.column_dimensions["B"].width = 130
    notas = [
        ("Metodología", None),
        ("", "Merma diaria = Sdo. Real − Sdo. Final, con Sdo. Final = Sdo. Inicial + Compras − Ventas + Ajustes (igual que el reporte del sistema)."),
        ("", "Merma mensual = suma de las mermas diarias = Sdo. Real del último día − (Sdo. Inicial del 1er día + Compras − Ventas + Ajustes). Merma % mensual = Merma mensual / Ventas del mes."),
        ("", "El % mensual es el indicador confiable: las mermas diarias se distorsionan cuando una pipa se registra un día distinto a su descarga (se ve como +30,000 L un día y −30,000 L al siguiente)."),
        ("", "Conciliación: primero se compara el total de compras del reporte de mermas contra el total de PIPAS DEL MES por estación y producto; después se busca cada pipa en las compras diarias, "
             "primero el mismo día y luego hasta ±N días (Parámetros). Una compra puede corresponder a la suma de 2+ pipas del mismo día."),
        ("", "Desfase = Fecha de registro en mermas − Fecha BOL. Positivo = se registró después de la fecha de la factura."),
        ("", "Costo promedio = promedio ponderado del CTO UNIT de las pipas del mes; Merma ($) = Merma (L) × costo promedio."),
        ("", "Registros manuales: lo que captures en la hoja 'Registros manuales' se suma a Compras (columna Compras manuales de Detalle diario) en la fecha indicada, "
             "y la pipa correspondiente pasa de NO REGISTRADA a REGISTRO MANUAL. Mermas, semáforos, conciliación y resumen se recalculan solos. "
             "La hoja 'Compras sin pipa' es la lista original del sistema; los registros manuales sin pipa se ven en la columna Resultado de 'Registros manuales'."),
        ("Equivalencias de nombres (PIPAS → estación oficial)", None),
    ] + [("", f"{k}  →  {v}") for k, v in ALIAS.items()] + [
        ("Archivos procesados", None), ("", f"Pipas: {os.path.basename(pipas_path)}")] + [("", f"Mermas: {os.path.basename(c)}") for c in sorted(csvs)] + [
        ("Avisos de calidad de datos", None)] + ([("", a) for a in avisos] or [("", "Sin avisos: los saldos son continuos y las sumas cuadran con los TOTALES de cada archivo.")])
    r = 4
    for a, b in notas:
        if b is None:
            x = ws_n.cell(r, 1, a); x.font = Font(name=F, size=11, bold=True, color=AZUL); r += 1; continue
        x = ws_n.cell(r, 2, "• " + b); x.font = Font(name=F, size=10); x.alignment = Alignment(wrap_text=True, vertical="top"); r += 1

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = {"Resumen": "1F3864", "Mermas mensual": "2F5597", "Detalle diario": "8EA9DB",
                                        "Conciliación pipas": "C00000", "Detalle pipas": "FF7C80", "Compras sin pipa": "FFC000",
                                        "Registros manuales": "5B9BD5", "Parámetros": "FFFF00", "Notas": "A6A6A6"}[ws.title]
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0
    wb.save(salida)
    return salida, periodo

def main(origen, salida=None):
    res = procesar(origen)
    diario = res[0]
    m = diario.Fecha.min()
    if salida is None:
        base = origen if os.path.isdir(origen) else os.path.dirname(os.path.abspath(origen))
        salida = os.path.join(base, f"Reporte_Mermas_{m:%Y-%m}.xlsx")
    construir_excel(salida, *res)
    print(f"✔ Reporte generado: {salida}")
    for a in res[4]: print("  ⚠", a)
    return salida

if __name__ == "__main__":
    origen = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
    salida = sys.argv[2] if len(sys.argv) > 2 else None
    main(origen, salida)
